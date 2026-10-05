import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_bab1_doc():
    doc = Document()

    # 1. Page Setup: A4, Margins: Left 4cm, Top 3cm, Right 3cm, Bottom 3cm (Standard Poliwangi)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)
    section.left_margin = Cm(4.0)
    section.right_margin = Cm(3.0)

    # Set normal style font to Times New Roman
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    # =========================================================================
    # HALAMAN COVER / JUDUL
    # =========================================================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.line_spacing = 1.15
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(18)
    
    r_title = p_title.add_run(
        "RANCANG BANGUN FITUR DIGITALISASI KOLEKSI DAN INTEGRASI QR CODE\n"
        "PADA SISTEM INFORMASI DI MUSEUM BLAMBANGAN BANYUWANGI\n\n"
        "LAPORAN MAGANG KERJA INDUSTRI"
    )
    r_title.bold = True
    r_title.font.size = Pt(13)

    # Add Poliwangi Logo
    if os.path.exists("logo_poliwangi.png"):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(12)
        p_logo.paragraph_format.space_after = Pt(12)
        run_logo = p_logo.add_run()
        run_logo.add_picture("logo_poliwangi.png", width=Cm(5.0))

    # Syarat kelulusan
    p_syarat = doc.add_paragraph()
    p_syarat.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_syarat.paragraph_format.line_spacing = 1.15
    p_syarat.paragraph_format.space_before = Pt(6)
    p_syarat.paragraph_format.space_after = Pt(18)
    r_syarat = p_syarat.add_run(
        "Magang Kerja Industri Dibuat dan Diajukan untuk Memenuhi Salah Satu Syarat Kelulusan\n"
        "Mata Kuliah Magang Kerja Industri Program Studi Sarjana Terapan Teknologi Rekayasa Perangkat Lunak\n"
        "Politeknik Negeri Banyuwangi"
    )
    r_syarat.font.size = Pt(10.5)

    # Penulis
    p_penulis = doc.add_paragraph()
    p_penulis.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_penulis.paragraph_format.line_spacing = 1.15
    p_penulis.paragraph_format.space_before = Pt(10)
    p_penulis.paragraph_format.space_after = Pt(24)
    r_penulis = p_penulis.add_run(
        "Oleh:\n"
        "SAMUEL RIVALDO SARAGIH\n"
        "NIM. 362358302156"
    )
    r_penulis.bold = True
    r_penulis.font.size = Pt(12)

    # Institusi
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.15
    p_inst.paragraph_format.space_before = Pt(18)
    p_inst.paragraph_format.space_after = Pt(0)
    r_inst = p_inst.add_run(
        "PROGRAM STUDI SARJANA TERAPAN\n"
        "TEKNOLOGI REKAYASA PERANGKAT LUNAK\n"
        "JURUSAN BISNIS DAN INFORMATIKA\n"
        "POLITEKNIK NEGERI BANYUWANGI\n"
        "2026"
    )
    r_inst.bold = True
    r_inst.font.size = Pt(12)

    doc.add_page_break()

    # =========================================================================
    # HALAMAN PENGESAHAN
    # =========================================================================
    p_peng_title = doc.add_paragraph()
    p_peng_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_peng_title.paragraph_format.line_spacing = 1.15
    p_peng_title.paragraph_format.space_after = Pt(18)
    r_peng_title = p_peng_title.add_run(
        "HALAMAN PENGESAHAN\n"
        "LAPORAN MAGANG KERJA INDUSTRI\n\n"
        "RANCANG BANGUN FITUR DIGITALISASI KOLEKSI DAN INTEGRASI QR CODE\n"
        "PADA SISTEM INFORMASI DI MUSEUM BLAMBANGAN BANYUWANGI"
    )
    r_peng_title.bold = True
    r_peng_title.font.size = Pt(12)

    p_peng_oleh = doc.add_paragraph()
    p_peng_oleh.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_peng_oleh.paragraph_format.space_after = Pt(18)
    r_peng_oleh = p_peng_oleh.add_run(
        "Oleh:\n"
        "SAMUEL RIVALDO SARAGIH\n"
        "NIM. 362358302156"
    )
    r_peng_oleh.bold = True
    r_peng_oleh.font.size = Pt(12)

    p_peng_tgl = doc.add_paragraph()
    p_peng_tgl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_peng_tgl.paragraph_format.space_after = Pt(20)
    r_peng_tgl = p_peng_tgl.add_run(
        "Telah diperiksa dan disetujui oleh pembimbing\n"
        "Pada tanggal: 15 Desember 2026."
    )
    r_peng_tgl.font.size = Pt(11)

    # Tabel Pembimbing
    t_pemb = doc.add_table(rows=2, cols=2)
    t_pemb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in t_pemb.rows:
        for cell in row.cells:
            cell.width = Cm(7.0)

    cell_p1 = t_pemb.cell(0, 0).paragraphs[0]
    cell_p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell_p1.paragraph_format.line_spacing = 1.15
    r = cell_p1.add_run("Pembimbing Lapang,\nMagang Kerja Industri,\n\n\n\n")
    r.font.size = Pt(11)
    r_name1 = cell_p1.add_run("Eko Ari Bawanto, S.Sn.\n")
    r_name1.bold = True
    r_name1.font.size = Pt(11)
    r_nip1 = cell_p1.add_run("NIP. 19940812 202406 1 002")
    r_nip1.font.size = Pt(11)

    cell_p2 = t_pemb.cell(0, 1).paragraphs[0]
    cell_p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell_p2.paragraph_format.line_spacing = 1.15
    r = cell_p2.add_run("Dosen Pembimbing,\nMagang Kerja Industri,\n\n\n\n")
    r.font.size = Pt(11)
    r_name2 = cell_p2.add_run("Alfin Hidayat, S.T., M.T.\n")
    r_name2.bold = True
    r_name2.font.size = Pt(11)
    r_nip2 = cell_p2.add_run("NIP. 19901005 201404 1 002")
    r_nip2.font.size = Pt(11)

    # Mengetahui
    p_mengetahui = doc.add_paragraph()
    p_mengetahui.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mengetahui.paragraph_format.space_before = Pt(20)
    p_mengetahui.paragraph_format.space_after = Pt(12)
    r_m = p_mengetahui.add_run("Mengetahui,")
    r_m.font.size = Pt(11)

    # Tabel Pejabat
    t_pej = doc.add_table(rows=1, cols=2)
    t_pej.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in t_pej.rows[0].cells:
        cell.width = Cm(7.0)

    cell_k1 = t_pej.cell(0, 0).paragraphs[0]
    cell_k1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell_k1.paragraph_format.line_spacing = 1.15
    r = cell_k1.add_run("Ketua Jurusan,\nBisnis dan Informatika,\n\n\n\n")
    r.font.size = Pt(11)
    r_k1 = cell_k1.add_run("Mohamad Dimyati Ayatullah, S.T., M.Kom.\n")
    r_k1.bold = True
    r_k1.font.size = Pt(11)
    r_k1_nip = cell_k1.add_run("NIPPPK. 19760122 202121 1 001")
    r_k1_nip.font.size = Pt(11)

    cell_k2 = t_pej.cell(0, 1).paragraphs[0]
    cell_k2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cell_k2.paragraph_format.line_spacing = 1.15
    r = cell_k2.add_run("Koordinator Program Studi,\nTeknologi Rekayasa Perangkat Lunak,\n\n\n\n")
    r.font.size = Pt(11)
    r_k2 = cell_k2.add_run("Dianni Yusuf, S.Kom., M.Kom.\n")
    r_k2.bold = True
    r_k2.font.size = Pt(11)
    r_k2_nip = cell_k2.add_run("NIPPPK. 19840305 202121 2 004")
    r_k2_nip.font.size = Pt(11)

    doc.add_page_break()

    # =========================================================================
    # KATA PENGANTAR
    # =========================================================================
    p_kp_title = doc.add_paragraph()
    p_kp_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kp_title.paragraph_format.space_after = Pt(18)
    r_kp_title = p_kp_title.add_run("KATA PENGANTAR")
    r_kp_title.bold = True
    r_kp_title.font.size = Pt(12)

    def add_body_p(text, indent=True, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(space_after)
        if indent:
            p.paragraph_format.first_line_indent = Cm(1.0)
        p.add_run(text)
        return p

    add_body_p(
        "Puji syukur kehadirat Allah SWT atas segala limpahan rahmat, hidayah, dan karunia-Nya, "
        "sehingga penulis dapat menyelesaikan Laporan Magang Kerja Industri (MKI) ini tepat pada waktunya. "
        "Laporan ini disusun dengan judul “Rancang Bangun Fitur Digitalisasi Koleksi dan Integrasi QR Code "
        "pada Sistem Informasi di Museum Blambangan Banyuwangi”."
    )

    add_body_p(
        "Penyusunan laporan ini dimaksudkan untuk memenuhi salah satu syarat kelulusan mata kuliah Magang Kerja "
        "Industri pada Program Studi Sarjana Terapan Teknologi Rekayasa Perangkat Lunak, Jurusan Bisnis dan Informatika, "
        "Politeknik Negeri Banyuwangi. Laporan ini merupakan hasil dari kegiatan praktik kerja, pengamatan, analisis "
        "kebutuhan sistem, serta pengembangan aplikasi yang penulis lakukan selama masa magang di Dinas Kebudayaan "
        "dan Pariwisata Kabupaten Banyuwangi pada unit Museum Blambangan."
    )

    add_body_p(
        "Penulis menyadari bahwa keberhasilan penyusunan laporan dan pelaksanaan magang ini tidak terlepas dari "
        "bantuan, bimbingan, arahan, serta dukungan dari berbagai pihak. Oleh karena itu, pada kesempatan ini penulis "
        "ingin menyampaikan rasa terima kasih yang sebesar-besarnya kepada:"
    )

    poin_terima_kasih = [
        "Bapak M. Shofiul Amin, S.T., M.T., selaku Direktur Politeknik Negeri Banyuwangi.",
        "Bapak Mohamad Dimyati Ayatullah, S.T., M.Kom., selaku Ketua Jurusan Bisnis dan Informatika Politeknik Negeri Banyuwangi.",
        "Ibu Dianni Yusuf, S.Kom., M.Kom., selaku Koordinator Program Studi Sarjana Terapan Teknologi Rekayasa Perangkat Lunak.",
        "Bapak Taufik Rohman, M.Si., selaku Kepala Dinas Kebudayaan dan Pariwisata Kabupaten Banyuwangi yang telah memberikan kesempatan dan fasilitas kepada penulis untuk melaksanakan kegiatan magang.",
        "Bapak Alfin Hidayat, S.T., M.T., selaku Dosen Pembimbing yang telah meluangkan waktu, tenaga, dan pikiran untuk membimbing penulis dalam penyusunan laporan ini.",
        "Bapak Eko Ari Bawanto, S.Sn., selaku Pembimbing Lapang di Dinas Kebudayaan dan Pariwisata yang senantiasa memberikan arahan teknis dan bimbingan selama penulis mengembangkan sistem di Museum Blambangan.",
        "Kedua orang tua dan keluarga tercinta yang senantiasa memberikan doa, dukungan moral, serta material yang tak terhingga.",
        "Rekan-rekan seperjuangan Magang Kerja Industri serta teman-teman mahasiswa Politeknik Negeri Banyuwangi yang saling memberikan semangat dan motivasi."
    ]

    for idx, poin in enumerate(poin_terima_kasih, 1):
        p_ptk = doc.add_paragraph()
        p_ptk.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ptk.paragraph_format.line_spacing = 1.5
        p_ptk.paragraph_format.left_indent = Cm(1.0)
        p_ptk.paragraph_format.first_line_indent = Cm(-0.6)
        p_ptk.paragraph_format.space_after = Pt(4)
        p_ptk.add_run(f"{idx}. {poin}")

    add_body_p(
        "Penulis menyadari bahwa laporan ini masih jauh dari kata sempurna, baik dari segi materi maupun teknik "
        "penyajiannya. Oleh karena itu, penulis sangat mengharapkan kritik dan saran yang membangun demi penyempurnaan "
        "laporan ini di masa mendatang. Akhir kata, penulis berharap semoga laporan ini dapat memberikan manfaat bagi "
        "pembaca, khususnya bagi pengembangan teknologi informasi di lingkungan kebudayaan dan permuseuman daerah.",
        indent=True,
        space_after=18
    )

    p_ttd = doc.add_paragraph()
    p_ttd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_ttd.paragraph_format.line_spacing = 1.15
    p_ttd.paragraph_format.space_after = Pt(0)
    r_ttd = p_ttd.add_run(
        "Banyuwangi, 15 Desember 2026\n\n\n\n"
        "Samuel Rivaldo Saragih\n"
        "NIM. 362358302156"
    )
    r_ttd.font.size = Pt(11)

    doc.add_page_break()

    # =========================================================================
    # BAB 1 PENDAHULUAN
    # =========================================================================
    p_bab1 = doc.add_paragraph()
    p_bab1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_bab1.paragraph_format.space_before = Pt(0)
    p_bab1.paragraph_format.space_after = Pt(18)
    r_bab1 = p_bab1.add_run("BAB 1\nPENDAHULUAN")
    r_bab1.bold = True
    r_bab1.font.size = Pt(12)

    def add_subheading(text, space_before=12, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(12)
        return p

    # 1.1 Latar Belakang Masalah
    add_subheading("1.1 Latar Belakang Masalah")

    add_body_p(
        "Politeknik Negeri Banyuwangi merupakan perguruan tinggi negeri vokasi yang berorientasi pada pengembangan "
        "sumber daya manusia profesional dengan penguasaan keahlian terapan (hard skill) dan kemampuan interpersonal "
        "(soft skill) yang kuat. Pada Program Studi Sarjana Terapan Teknologi Rekayasa Perangkat Lunak (TRPL), proses "
        "pembelajaran diterapkan melalui pendekatan Project Based Learning (PBL) yang menekankan penyelesaian masalah "
        "nyata melalui proyek perangkat lunak terstruktur. Pendekatan ini dirancang agar mahasiswa mampu menerapkan "
        "konsep rekayasa perangkat lunak secara langsung pada permasalahan teknis yang relevan dengan industri. Hal ini "
        "sejalan dengan pendapat Thomas (2000) yang menyatakan bahwa “Pembelajaran berbasis proyek mendorong pemahaman "
        "yang lebih mendalam dan membantu mahasiswa menghubungkan materi akademik dengan praktik profesional di dunia nyata”. "
        "Dengan demikian pembelajaran berbasis proyek tersebut menjadi dasar penting dalam pelaksanaan Magang Kerja Industri (MKI), "
        "karena mahasiswa diarahkan untuk menerapkan pengetahuan dan pengalaman belajar mereka secara langsung dalam "
        "lingkungan kerja sesungguhnya."
    )

    add_body_p(
        "Dinas Kebudayaan dan Pariwisata Kabupaten Banyuwangi, yang berlokasi di Jl. Jenderal Ahmad Yani No. 78, "
        "memiliki tanggung jawab dalam penyelenggaraan urusan pemerintah daerah di bidang kebudayaan dan kepariwisataan. "
        "Berdasarkan dokumen Rencana Strategis (Renstra), salah satu fokus utama Disbudpar adalah meningkatkan pelestarian "
        "cagar budaya, kesenian daerah, serta memperkuat pengelolaan permuseuman sebagai bagian integral dari upaya "
        "pelestarian peradaban dan pengembangan pariwisata berbasis budaya lokal (heritage tourism). Salah satu unit "
        "strategis yang dikelola oleh dinas adalah Museum Blambangan Banyuwangi, yang menyimpan lebih dari 4.000 koleksi "
        "benda cagar budaya mencakup peninggalan masa prasejarah, era Kerajaan Majapahit dan Kerajaan Blambangan, masa kolonial, "
        "hingga keragaman etnografi masyarakat Osing. Keberadaan data dan akses informasi yang akurat sangat krusial dalam "
        "menunjang tugas pelestarian tersebut. Sebagaimana yang ditegaskan oleh Koentjaraningrat (1985), bahwa “pelestarian "
        "budaya tidak dapat dilakukan tanpa data yang akurat”. Informasi cagar budaya menjadi basis strategis bagi pemerintah "
        "daerah dalam melaksanakan edukasi publik, menyusun program kerja kebudayaan, serta menjaga keberlanjutan warisan sejarah lokal."
    )

    add_body_p(
        "Sebagai sarana edukasi masyarakat dan daya tarik wisata sejarah, Museum Blambangan memerlukan sistem penyajian "
        "informasi yang modern dan mudah diakses. Namun, hasil observasi langsung dan evaluasi lapangan menunjukkan bahwa "
        "pengelolaan informasi koleksi serta pelayanan tiket di Museum Blambangan masih menghadapi sejumlah kendala mendasar. "
        "Pada aspek penyajian koleksi, media informasi yang tersedia masih mengandalkan label fisik konvensional berbahan kertas "
        "atau akrilik yang ditempel di samping etalase pameran. Label fisik ini memiliki dimensi yang sangat terbatas, sehingga hanya "
        "mampu memuat ringkasan nama dan asal benda tanpa dapat memaparkan latar belakang sejarah, filosofi, serta dokumentasi visual "
        "yang mendalam. Selain itu, ketiadaan dukungan multibahasa menjadi hambatan komunikasi bagi wisatawan mancanegara yang berkunjung. "
        "Sistem pameran konvensional tersebut juga belum menyediakan fasilitas audio pemandu (audio guide), sehingga membatasi aksesibilitas "
        "bagi pengunjung penyandang disabilitas (khususnya tunanetra) serta mengurangi daya tarik interaktif bagi generasi muda."
    )

    add_body_p(
        "Di sisi lain, pada aspek pelayanan kunjungan (ticketing), pencatatan pengunjung dan retribusi tiket masuk masih dikelola "
        "secara semi-manual menggunakan karcis fisik. Kondisi ini menimbulkan potensi antrean di loket pada saat jam kunjungan padat, "
        "menyulitkan pengawasan kapasitas kuota harian per sesi kunjungan, serta rentan terhadap risiko kesalahan manusia (human error) "
        "dalam rekapitulasi data pendapatan kas Pendapatan Asli Daerah (PAD). Selain itu, belum adanya mekanisme validasi tiket otomatis "
        "di pintu masuk gerbang museum menyebabkan petugas loket harus memeriksa karcis secara visual satu per satu, yang memperlambat "
        "proses masuk dan berisiko terjadinya duplikasi pemakaian tiket."
    )

    add_body_p(
        "Berdasarkan permasalahan nyata tersebut, modernisasi sistem layanan museum menjadi kebutuhan mendesak demi menjamin penyajian "
        "koleksi yang inklusif, terstandar, dan berdaya saing global, sekaligus mewujudkan tata kelola tiket yang transparan dan akuntabel. "
        "Teknologi Quick Response Code (QR Code) menawarkan solusi yang tepat, efektif, dan berbiaya rendah karena memiliki kapasitas simpan "
        "data yang tinggi, kemampuan koreksi kesalahan (error correction), serta dapat dipindai secara instan menggunakan kamera telepon pintar "
        "(smartphone) pengunjung tanpa mewajibkan instalasi aplikasi khusus dari luar."
    )

    add_body_p(
        "Melalui pelaksanaan MKI, penulis berfokus pada rancang bangun fitur digitalisasi koleksi cagar budaya (Musewangi) dan sistem "
        "pemesanan tiket elektronik (Booking Museum) dengan integrasi teknologi QR Code pada Sistem Informasi Museum Blambangan Banyuwangi. "
        "Fitur Musewangi menyediakan pembuatan dan pengelolaan QR Code koleksi, pemindaian QR menggunakan kamera web, penyajian katalog "
        "digital komprehensif, dukungan informasi dalam tiga bahasa (Bahasa Indonesia, English, dan Basa Osing), serta fitur pemutar audio "
        "deskripsi narasi suara. Sementara itu, fitur Booking Museum mencakup pemesanan tiket dengan pembagian sesi kuota, pembayaran resmi "
        "kas daerah berbasis QRIS dan transfer Bank Jatim, penerbitan E-Tiket ber-QR Code yang dapat diunduh dalam format PDF, serta modul "
        "validasi gerbang masuk (gate check-in scanner) untuk petugas loket. Transformasi digital ini diharapkan mampu memperkuat "
        "infrastruktur pelayanan Dinas Kebudayaan dan Pariwisata Kabupaten Banyuwangi dalam mewujudkan layanan museum berstandar modern."
    )

    # 1.2 Rumusan Masalah
    add_subheading("1.2 Rumusan Masalah")
    
    rumusan = [
        "Bagaimana merancang dan membangun fitur digitalisasi koleksi cagar budaya (Musewangi) berbasis teknologi QR Code pada Sistem Informasi Museum Blambangan Banyuwangi?",
        "Bagaimana mengimplementasikan penyajian informasi koleksi digital yang mendukung format multibahasa (Bahasa Indonesia, English, Basa Osing) serta fitur audio deskripsi suara interaktif?",
        "Bagaimana merancang dan mengintegrasikan sistem tiket kunjungan elektronik (Booking Museum) yang memfasilitasi pembuatan QR Code tiket, pengunduhan berkas PDF, serta validasi tiket di pintu masuk secara real-time?"
    ]
    for idx, r_item in enumerate(rumusan, 1):
        p_rm = doc.add_paragraph()
        p_rm.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_rm.paragraph_format.line_spacing = 1.5
        p_rm.paragraph_format.left_indent = Cm(1.0)
        p_rm.paragraph_format.first_line_indent = Cm(-0.6)
        p_rm.paragraph_format.space_after = Pt(4)
        p_rm.add_run(f"{idx}. {r_item}")

    # 1.3 Tujuan
    add_subheading("1.3 Tujuan")

    tujuan = [
        "Mengimplementasikan fitur digitalisasi koleksi (Musewangi) pada Museum Blambangan melalui pembuatan QR Code unik untuk setiap benda cagar budaya serta pemindaian QR berbasis kamera peramban web.",
        "Menyajikan antarmuka detail koleksi sejarah yang informatif, responsif, dan inklusif dengan dukungan multibahasa (Indonesia, Inggris, Osing) serta fitur pemutar audio deskripsi suara bagi kenyamanan pengunjung dan penyandang disabilitas.",
        "Membangun modul pemesanan tiket kunjungan daring (Booking Museum) terintegrasi kas daerah (PAD) yang menerbitkan E-Tiket ber-QR Code dengan fitur unduh berkas PDF dan modul pemindaian validasi check-in di pintu gerbang museum."
    ]
    for idx, t_item in enumerate(tujuan, 1):
        p_tj = doc.add_paragraph()
        p_tj.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_tj.paragraph_format.line_spacing = 1.5
        p_tj.paragraph_format.left_indent = Cm(1.0)
        p_tj.paragraph_format.first_line_indent = Cm(-0.6)
        p_tj.paragraph_format.space_after = Pt(4)
        p_tj.add_run(f"{idx}. {t_item}")

    # 1.4 Manfaat
    add_subheading("1.4 Manfaat")

    # 1.4.1 Disbudpar
    p_m1 = doc.add_paragraph()
    p_m1.paragraph_format.left_indent = Cm(0.5)
    p_m1.paragraph_format.space_before = Pt(4)
    p_m1.paragraph_format.space_after = Pt(2)
    p_m1.paragraph_format.keep_with_next = True
    r_m1 = p_m1.add_run("1.4.1 Manfaat Bagi Dinas Kebudayaan Dan Pariwisata")
    r_m1.bold = True
    r_m1.font.size = Pt(12)

    manfaat_disbudpar = [
        "Meningkatkan mutu edukasi dan daya tarik Museum Blambangan melalui sarana informasi digital interaktif, multibahasa, dan ramah aksesibilitas.",
        "Meningkatkan efisiensi kerja, transparansi, serta akuntabilitas pencatatan pendapatan retribusi tiket daerah (PAD) melalui sistem pembayaran digital dan reservasi terstruktur.",
        "Membantu petugas loket dalam mempercepat proses verifikasi pengunjung di pintu masuk serta mencegah terjadinya kecurangan tiket ganda melalui pemindaian QR Code."
    ]
    for idx, mb in enumerate(manfaat_disbudpar, 1):
        p_mb = doc.add_paragraph()
        p_mb.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_mb.paragraph_format.line_spacing = 1.5
        p_mb.paragraph_format.left_indent = Cm(1.5)
        p_mb.paragraph_format.first_line_indent = Cm(-0.6)
        p_mb.paragraph_format.space_after = Pt(3)
        p_mb.add_run(f"{idx}. {mb}")

    # 1.4.2 Poliwangi
    p_m2 = doc.add_paragraph()
    p_m2.paragraph_format.left_indent = Cm(0.5)
    p_m2.paragraph_format.space_before = Pt(6)
    p_m2.paragraph_format.space_after = Pt(2)
    p_m2.paragraph_format.keep_with_next = True
    r_m2 = p_m2.add_run("1.4.2 Manfaat Bagi Politeknik Negeri Banyuwangi")
    r_m2.bold = True
    r_m2.font.size = Pt(12)

    manfaat_poliwangi = [
        "Memperkuat implementasi kurikulum vokasi berbasis proyek (Project Based Learning) dan keterkaitan kerja sama (link and match) dengan instansi pemerintah daerah.",
        "Memberikan umpan balik akademik yang nyata bagi Program Studi TRPL mengenai penerapan rekayasa perangkat lunak modern pada sektor kebudayaan dan pariwisata.",
        "Menambah khazanah referensi ilmiah dan karya terapan mahasiswa di bidang integrasi teknologi QR Code, pemrosesan media audio, dan sistem e-ticketing."
    ]
    for idx, mp in enumerate(manfaat_poliwangi, 1):
        p_mp = doc.add_paragraph()
        p_mp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_mp.paragraph_format.line_spacing = 1.5
        p_mp.paragraph_format.left_indent = Cm(1.5)
        p_mp.paragraph_format.first_line_indent = Cm(-0.6)
        p_mp.paragraph_format.space_after = Pt(3)
        p_mp.add_run(f"{idx}. {mp}")

    # 1.4.3 Mahasiswa
    p_m3 = doc.add_paragraph()
    p_m3.paragraph_format.left_indent = Cm(0.5)
    p_m3.paragraph_format.space_before = Pt(6)
    p_m3.paragraph_format.space_after = Pt(2)
    p_m3.paragraph_format.keep_with_next = True
    r_m3 = p_m3.add_run("1.4.3 Manfaat Bagi Mahasiswa")
    r_m3.bold = True
    r_m3.font.size = Pt(12)

    manfaat_mhs = [
        "Meningkatkan kompetensi teknis dalam perancangan aplikasi web modern berbasis React 19, TypeScript, Tailwind CSS, engine QR Code, manipulasi canvas, serta Web Speech API.",
        "Memperoleh pengalaman nyata dalam menganalisis kebutuhan operasional instansi pemerintah dan mengimplementasikan solusi rekayasa perangkat lunak yang solutif.",
        "Melatih profesionalisme, kemampuan komunikasi interpersonal, dan etika kerja dalam berinteraksi dengan aparatur sipil negara dan masyarakat luas."
    ]
    for idx, mm in enumerate(manfaat_mhs, 1):
        p_mm = doc.add_paragraph()
        p_mm.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_mm.paragraph_format.line_spacing = 1.5
        p_mm.paragraph_format.left_indent = Cm(1.5)
        p_mm.paragraph_format.first_line_indent = Cm(-0.6)
        p_mm.paragraph_format.space_after = Pt(3)
        p_mm.add_run(f"{idx}. {mm}")

    # 1.5 Batasan Masalah
    add_subheading("1.5 Batasan Masalah")

    add_body_p(
        "Untuk menjaga fokus pembahasan agar terarah dan sesuai dengan sasaran penugasan Magang Kerja Industri, "
        "penulis menetapkan batasan masalah sistem sebagai berikut:"
    )

    batasan = [
        "Sistem dikembangkan berbasis aplikasi web responsif menggunakan pustaka React 19, bahasa pemrograman TypeScript dengan pengecekan tipe data ketat (Strict Mode), dan kerangka kerja desain Tailwind CSS.",
        "Pengembangan modul digitalisasi koleksi (Musewangi) dibatasi pada pembuatan QR Code koleksi, pemindaian QR Code menggunakan kamera peramban web (jsQR API), penyajian informasi detail artefak, penyediaan opsi tiga bahasa (Bahasa Indonesia, English, dan Basa Osing), serta fitur pemutar audio deskripsi narasi suara.",
        "Pengembangan modul tiket kunjungan (Booking Museum) difokuskan pada alur pemesanan tiket dengan batasan kuota per sesi kunjungan, metode pembayaran resmi kas daerah (QRIS / Bank Jatim), penerbitan E-Tiket ber-QR Code, pengunduhan tiket berformat PDF (jspdf dan html2canvas), serta validasi check-in gerbang pintu masuk oleh petugas loket.",
        "Manajemen data dan persistensi status pemesanan dikelola menggunakan React Context API terpusat yang tersinkronisasi dengan penyimpanan lokal terstruktur (structured local storage) untuk kebutuhan simulasi operasional dinas dan demonstrasi sistem."
    ]
    for idx, b_item in enumerate(batasan, 1):
        p_bt = doc.add_paragraph()
        p_bt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_bt.paragraph_format.line_spacing = 1.5
        p_bt.paragraph_format.left_indent = Cm(1.0)
        p_bt.paragraph_format.first_line_indent = Cm(-0.6)
        p_bt.paragraph_format.space_after = Pt(4)
        p_bt.add_run(f"{idx}. {b_item}")

    output_filename = "LAPORAN_MAGANG_BAB_1_SAMUEL_RIVALDO.docx"
    doc.save(output_filename)
    print(f"Berhasil membuat dokumen: {output_filename}")

if __name__ == "__main__":
    create_bab1_doc()
